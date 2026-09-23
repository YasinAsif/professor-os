from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION_START
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


SRC = Path(r"E:\FYP_YASIN\ProfessorOS_FYP_Report_Formatted.docx")
OUT = Path(r"E:\FYP_YASIN\ProfessorOS_FYP_Report_30pct.docx")


def set_font(run, size=11, bold=None, italic=None, color=None):
    run.font.name = "Times New Roman"
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), "Times New Roman")
    rfonts.set(qn("w:hAnsi"), "Times New Roman")
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = RGBColor(*color)


def clear_paragraph(p):
    for child in list(p._p):
        if child.tag != qn("w:pPr"):
            p._p.remove(child)


def set_paragraph_text(p, text, size=11, bold=False, italic=False, color=None, align=None):
    clear_paragraph(p)
    if align is not None:
        p.alignment = align
    r = p.add_run(text)
    set_font(r, size=size, bold=bold, italic=italic, color=color)
    return p


def add_page_number(paragraph):
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr)
    run._r.append(fld_char2)
    set_font(run, size=9, color=(90, 90, 90))


def add_bottom_border(p, color="B7C9D6", size="6"):
    ppr = p._p.get_or_add_pPr()
    pbdr = ppr.find(qn("w:pBdr"))
    if pbdr is None:
        pbdr = OxmlElement("w:pBdr")
        ppr.append(pbdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)
    pbdr.append(bottom)


def replace_next_paragraph(doc, heading, new_text):
    paras = doc.paragraphs
    for i, p in enumerate(paras):
        if " ".join(p.text.split()) == heading:
            for q in paras[i + 1:]:
                if q.text.strip():
                    set_paragraph_text(q, new_text, size=11)
                    q.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                    return


def add_header_footer(doc):
    for si, section in enumerate(doc.sections):
        section.header_distance = Inches(0.45)
        section.footer_distance = Inches(0.45)
        section.different_first_page_header_footer = (si == 0)

        header = section.header
        p = header.paragraphs[0]
        clear_paragraph(p)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run("ProfessorOS  |  30% Implementation Report")
        set_font(r, size=9, bold=True, color=(31, 78, 121))
        add_bottom_border(p)

        footer = section.footer
        fp = footer.paragraphs[0]
        clear_paragraph(fp)
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = fp.add_run("Shifa Tameer-e-Millat University  |  Page ")
        set_font(r, size=9, color=(90, 90, 90))
        add_page_number(fp)


def improve_title_pages(doc):
    for i in (1, 15):
        p = doc.paragraphs[i]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(30)
        p.paragraph_format.space_after = Pt(16)
        for r in p.runs:
            set_font(r, size=28, bold=True, color=(31, 78, 121))
    for i in (2, 16):
        p = doc.paragraphs[i]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(16)
        for r in p.runs:
            set_font(r, size=16, bold=False, color=(31, 78, 121))
    for i in range(3, min(24, len(doc.paragraphs))):
        p = doc.paragraphs[i]
        if p.text.strip():
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(8)
            for r in p.runs:
                set_font(r, size=13)
    # Give the title-page sections more breathing room without changing the report page geometry.
    for section in doc.sections[:2]:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)


def fill_caption_lists(doc):
    tables = [p.text.strip() for p in doc.paragraphs if p.text.strip().startswith("Table ") and not p.text.strip().startswith("Table captions have been standardized")]
    figures = [p.text.strip() for p in doc.paragraphs if p.text.strip().startswith("Figure ")]
    active = None
    for p in doc.paragraphs:
        t = " ".join(p.text.split())
        if t == "LIST OF TABLES":
            active = "tables"
        elif t == "LIST OF FIGURES":
            active = "figures"
        elif active == "tables" and (t.startswith("Table captions have been standardized") or t.startswith("Figure captions have been standardized")):
            lines = "List of Tables\n" + "\n".join(tables)
            set_paragraph_text(p, lines, size=10)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.line_spacing = 1.0
            active = None
        elif active == "figures" and (t.startswith("Table captions have been standardized") or t.startswith("Figure captions have been standardized")):
            lines = "List of Figures\n" + "\n".join(figures)
            set_paragraph_text(p, lines, size=10)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.line_spacing = 1.0
            active = None


def revise_status_content(doc):
    replace_next_paragraph(doc, "CHAPTER 5 IMPLEMENTATION AND INTEGRATION",
        "This chapter reports the implemented 30% milestone rather than claiming completion of the full ProfessorOS scope. The completed increment establishes the working foundation: FastAPI services, PostgreSQL and SQLAlchemy data models, Argon2id and JWT security, role-based access control, course and roster onboarding, HEC CLO and rubric groundwork, Flutter web/mobile application structure, analytics endpoints, Docker-based development configuration, and an initial automated test suite. AI grading, explainable AI, smart proctoring, code execution, the AI teaching assistant, full end-to-end integration, and production hardening are explicitly deferred to the remaining 70% implementation.")
    replace_next_paragraph(doc, "5.2 Repository Structure and Configuration Management",
        "The 30% implementation is organized into backend, frontend, selenium_tests, evaluation_artifacts, and documentation areas. The backend contains FastAPI API routes, security and configuration code, SQLAlchemy models, schemas, services, tasks, and migrations. The frontend contains the Flutter application, shared design tokens, feature modules, web entry point, and Android build structure. This repository organization supports incremental delivery; final release branching, tagging, and complete deployment packaging will be finalized after the remaining modules are implemented.")
    replace_next_paragraph(doc, "5.3 Implementation of Major Modules",
        "The 30% increment completes the core foundation needed for later modules. Implemented work includes multi-role authentication with Argon2id and JWT, user approval and access rules, course creation and student onboarding, join-code and CSV roster workflows, assignment and rubric data structures, CLO mapping groundwork, analytics API foundations, the Flutter application shell, and the Marginalia visual system. The remaining 70% will implement and integrate AI grading, XAI explanations, submission processing, smart proctoring, code execution, RAG-based teaching assistance, richer analytics, and final production workflows.")
    replace_next_paragraph(doc, "5.6 User Interface Implementation",
        "The 30% deliverable includes the Flutter web/mobile application shell, adaptive navigation, authentication screens, course and administration flows, reusable components, and the Marginalia paper-and-ink design system. Final workflow coverage, implementation screenshots for every role, accessibility evidence, and polished submission, grading, proctoring, and teaching-assistant screens will be completed during the remaining 70% implementation.")
    replace_next_paragraph(doc, "5.7 System Integration",
        "At the 30% stage, the foundation integration is in place between the Flutter client, FastAPI services, authentication layer, course-management APIs, database models, and development deployment configuration. Full integration across AI grading, XAI, proctoring, Judge0, FAISS/RAG, notification flows, and production observability is planned for the remaining 70% and is not claimed as complete in this report.")
    replace_next_paragraph(doc, "5.8 Coding Quality, Build Automation, and Technical Documentation",
        "The repository includes modular backend and frontend structures, typed request/response schemas, explicit role guards, API documentation through FastAPI, Docker Compose configuration, and separate backend and Selenium test areas. Complete CI evidence, static-analysis baselines, release tagging, and final technical documentation will be added after the remaining implementation modules stabilize.")
    replace_next_paragraph(doc, "5.9 Release Preparation and Deployment Readiness",
        "The 30% milestone is a development increment and is not presented as a production release. Docker-based local deployment configuration, a Flutter Android build artifact, and the deployed-test project structure provide an initial delivery foundation. Production smoke testing, monitoring, backup and rollback evidence, signed release packaging, and operational handover are reserved for the remaining 70%.")

    replace_next_paragraph(doc, "CHAPTER 6 VERIFICATION, VALIDATION, TESTING, AND EVALUATION",
        "Testing at the 30% milestone is limited to the implemented foundation and its accessible workflows. The repository contains backend pytest tests for approval idempotency, registration policy, assignment access, student access, and dashboard behavior, together with Selenium tests for smoke loading, authentication outcomes, admin navigation, and course-management controls. Tests requiring deployed credentials are intentionally skipped until the target environment is configured. Full performance, security, AI-quality, proctoring, end-to-end, and user-acceptance evaluation will be completed after the remaining 70% is implemented.")
    testing = {
        "6.1 Evaluation Strategy and Test Environment": "The 30% evaluation strategy combines code-level pytest coverage with browser-level Selenium smoke and navigation tests. The browser suite targets the deployed Flutter Web surface; authenticated cases require APP_URL, ADMIN_EMAIL, and ADMIN_PASSWORD. This stage validates the foundation only, not the complete proposed system.",
        "6.2 Unit Testing": "Backend unit tests are present for approval-state idempotency, registration email policy, assignment access, student access, and student dashboard behavior. These tests exercise service and authorization rules that support the 30% foundation. Additional unit tests for AI grading, XAI, proctoring, code execution, and RAG will be added with those modules.",
        "6.3 Functional and Business-Rule Testing": "The current functional scope covers authentication outcomes, account approval, role restrictions, course-management access, and the initial course and roster workflows. Boundary cases such as empty credentials and invalid passwords are represented in the Selenium authentication tests. Assignment grading and dispute rules remain part of the next implementation increment.",
        "6.4 Integration Testing": "Initial integration is checked at the client-to-API and API-to-data-model boundaries for the implemented foundation. A complete integration run across all proposed services is deferred because the AI, XAI, Judge0, proctoring, vector-search, and notification components are not part of the 30% completed increment.",
        "6.5 System and End-to-End Testing": "The 30% system-level check is limited to reaching the deployed login surface, rendering a non-empty document, and exercising available authenticated administration and course-management controls. Full student submission-to-grade and professor review journeys will be tested after the remaining modules are integrated.",
        "6.6 Non-Functional Requirement Testing": "At this stage, non-functional checks are limited to basic application reachability, input handling, role protection, and responsive Flutter structure. Formal load, latency, uptime, scalability, and offline-behavior measurements are planned for the final implementation.",
        "6.7 Security and Privacy Testing": "The implemented foundation is designed around Argon2id password hashing, JWT authentication, role guards, approval rules, and structured HTTP errors. The 30% testing scope checks access-control behavior and invalid authentication paths. Full penetration testing, privacy review, and production secret/configuration audit remain future work.",
        "6.8 Usability and User Acceptance Validation": "The 30% usability review covers the implemented authentication, administration, course-management, and adaptive navigation surfaces using the Marginalia design system. Formal user acceptance sessions with professors, teaching assistants, and students will be conducted once the full assessment lifecycle is available.",
        "6.9 Results and Defect Summary": "This milestone report records the test assets and scope, but it does not claim final pass-rate or defect-closure evidence for the whole system. Authenticated browser tests depend on deployment credentials and environment availability. A complete results table and defect log will be added after the remaining 70% test cycle.",
        "6.10 Requirements and Objectives Validation": "The 30% evidence supports the foundation requirements for authentication, role control, course setup, onboarding, and initial analytics/UI structure. Requirements for automated grading, explainability, proctoring, code execution, RAG assistance, and full deployment are intentionally marked as pending implementation and validation.",
        "6.11 Discussion of Evaluation Findings": "The evaluation confirms that the project has moved from design into a working foundation increment. The main finding is that core architecture and access-controlled workflows are testable, while the highest-risk assessment features remain in the remaining 70% work. This scope boundary keeps the report evidence aligned with the 30% submission.",
    }
    for h, text in testing.items():
        replace_next_paragraph(doc, h, text)

    # Replace remaining generic placeholder text outside the 30% implementation chapters.
    for p in doc.paragraphs:
        if "[Content not available in current project documents - to be completed during subsequent development phases.]" in p.text:
            set_paragraph_text(p, "This activity is planned for the remaining 70% implementation and will be completed when the related module and evidence become available.", size=11, italic=True, color=(90, 90, 90))


def main():
    doc = Document(str(SRC))
    improve_title_pages(doc)
    fill_caption_lists(doc)
    revise_status_content(doc)
    add_header_footer(doc)
    # Keep field results refreshable when opened in Word.
    settings = doc.settings.element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")
    doc.core_properties.subject = "30% Implementation Milestone Report"
    doc.save(str(OUT))
    print(OUT)


if __name__ == "__main__":
    main()
